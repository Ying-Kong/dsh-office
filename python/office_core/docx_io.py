"""Low level .docx access.

Only this module touches python-docx objects. Everything above it works with
plain dictionaries, which keeps the op interpreter and the verifier testable
without a live Word document.
"""

from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.document import Document as _Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .errors import UnsupportedDocument

# Parts whose bytes decide the document's content. Relationships are excluded
# on purpose: python-docx assigns rIds in a stable order, but the mapping is an
# implementation detail, not content.
_CONTENT_PARTS = ("word/document.xml", "word/styles.xml", "word/numbering.xml")


def load_docx(path: str | Path) -> _Document:
    """Open a .docx and map failures onto core error types."""
    path = Path(path)
    if not path.is_file():
        raise UnsupportedDocument(f"not a file: {path}", path=str(path))
    if path.suffix.lower() == ".doc":
        raise UnsupportedDocument("legacy .doc is not supported; convert to .docx first", path=str(path))
    try:
        return Document(str(path))
    except zipfile.BadZipFile as exc:  # pragma: no cover - defensive
        raise UnsupportedDocument(f"not an OOXML package: {path}", path=str(path)) from exc


def iter_blocks(doc: _Document) -> Iterator[Paragraph | Table]:
    """Yield body paragraphs and tables in document order."""
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def logical_text(paragraph: Paragraph) -> str:
    """Paragraph text with Word's run splitting undone.

    Word splits a sentence into several ``w:r`` whenever formatting, spelling
    state or revision marks change mid-sentence. Matching user-visible text
    against the raw runs therefore fails. This joins every ``w:t`` in order and
    renders ``w:br``/``w:tab`` the way a reader would see them.
    """
    out: list[str] = []
    for node in paragraph._p.iter():
        tag = node.tag
        if tag == qn("w:t"):
            out.append(node.text or "")
        elif tag == qn("w:tab"):
            out.append("\t")
        elif tag in (qn("w:br"), qn("w:cr")):
            out.append("\n")
    return "".join(out)


def runs_of(paragraph: Paragraph) -> list:
    return list(paragraph.runs)


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """Replace a paragraph's text, keeping the first run's formatting.

    This is the only write path for existing text. It deliberately does not
    touch ``pPr`` (spacing, borders, page breaks) so template layout survives,
    and it removes the surplus runs instead of leaving empty ones behind.
    """
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run._element.getparent().remove(run._element)


def cell_text(cell) -> str:
    parts = []
    for paragraph in cell.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    return "\n".join(parts)


def set_cell_text(cell, text: str) -> None:
    """Write a cell's text into its first paragraph, keeping its formatting."""
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.add_paragraph(text)
        return
    replace_paragraph_text(paragraphs[0], text)
    for extra in paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def image_count(doc: _Document) -> int:
    return len(doc.inline_shapes)


def content_fingerprint(path: str | Path) -> str:
    """Hash the content parts of a package.

    Two renders of the same template with the same ops must produce the same
    fingerprint. Byte identity of the whole .zip is *not* guaranteed (zip entry
    timestamps differ); see docs/limitations.md.
    """
    path = Path(path)
    digest = hashlib.sha256()
    with zipfile.ZipFile(path) as package:
        names = sorted(n for n in package.namelist() if n in _CONTENT_PARTS or n.startswith("word/media/"))
        for name in names:
            digest.update(name.encode())
            digest.update(hashlib.sha256(package.read(name)).digest())
    return digest.hexdigest()
