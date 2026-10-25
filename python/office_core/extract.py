"""Turn a .docx into a flat, reviewable fact table.

The panel is generated from this structure; the verifier diffs against it.
Nothing here interprets intent - it only reports what the tree contains, plus
the handful of heuristics that guess which empty spots are meant to be filled.
"""

from __future__ import annotations

from pathlib import Path

from docx.oxml.ns import qn

from .docx_io import cell_text, content_fingerprint, image_count, load_docx, logical_text
from .index import DocIndex
from .schemas import FACTS_SCHEMA
from .styles import effective_paragraph_format, effective_run_format


def _page_info(doc) -> dict:
    section = doc.sections[0]
    return {
        "width_twip": int(section.page_width.twips) if section.page_width else None,
        "height_twip": int(section.page_height.twips) if section.page_height else None,
        "margin_twip": {
            "top": int(section.top_margin.twips) if section.top_margin else None,
            "bottom": int(section.bottom_margin.twips) if section.bottom_margin else None,
            "left": int(section.left_margin.twips) if section.left_margin else None,
            "right": int(section.right_margin.twips) if section.right_margin else None,
        },
    }


def _paragraph_facts(doc, pid: str, paragraph) -> dict:
    runs = paragraph.runs
    fmt = effective_run_format(doc, paragraph, runs[0]) if runs else {
        "sz_pt": None, "font_ascii": None, "font_east_asia": None,
        "bold": None, "italic": None, "style_name": paragraph.style.name if paragraph.style else None,
    }
    pfmt = effective_paragraph_format(doc, paragraph)
    has_drawing = paragraph._p.find(".//" + qn("w:drawing")) is not None
    return {
        "id": pid,
        "text": logical_text(paragraph),
        "sz_pt": fmt["sz_pt"],
        "font_ascii": fmt["font_ascii"],
        "font_east_asia": fmt["font_east_asia"],
        "bold": fmt["bold"],
        "italic": fmt["italic"],
        "style_name": fmt["style_name"],
        "align": pfmt["align"],
        "has_image": has_drawing,
    }


def _table_facts(index: DocIndex) -> list[dict]:
    tables: dict[str, dict] = {}
    for cid in index.cell_order:
        table_id, rest = cid.split("r", 1)
        row_index = int(rest.split("c", 1)[0])
        table = tables.setdefault(table_id, {"id": table_id, "rows": {}})
        table["rows"].setdefault(row_index, []).append({
            "id": cid,
            "text": cell_text(index.cells[cid]).strip(),
            "empty": cell_text(index.cells[cid]).strip() == "",
        })
    # normalise insertion order back to row order
    return [
        {"id": t["id"], "rows": [t["rows"][r] for r in sorted(t["rows"])]}
        for t in tables.values()
    ]


def _detect_slots(index: DocIndex) -> list[dict]:
    slots: list[dict] = []
    counter = 0

    def add(**kwargs):
        nonlocal counter
        counter += 1
        slots.append({"id": f"s{counter}", **kwargs})

    # 1) empty cell to the right of a short label - the report-form pattern
    for cid in index.cell_order:
        text = cell_text(index.cells[cid]).strip()
        if not text or len(text) > 14 or "\n" in text:
            continue
        neighbour = index.next_cell_id(cid)
        if neighbour is not None and cell_text(index.cells[neighbour]).strip() == "":
            add(kind="cell", label=text, current="",
                target={"by": "cell-id", "id": neighbour},
                locator={"by": "cell-label", "label": text},
                confidence=0.9, source="heuristic:label-right-cell")

    # 2) inline date line, e.g. 日期：    2025年  月  日
    for pid, paragraph in index.paragraphs.items():
        text = logical_text(paragraph)
        if "日期" in text and all(token in text for token in ("年", "月", "日")):
            add(kind="date-inline", label="日期", current=text,
                target={"by": "paragraph-id", "id": pid},
                locator={"by": "paragraph-id", "id": pid},
                confidence=0.75, source="heuristic:date-line")

    # 3) explicit blanks
    for pid, paragraph in index.paragraphs.items():
        text = logical_text(paragraph)
        if "___" in text:
            add(kind="blank", label=text.strip()[:20], current=text,
                target={"by": "paragraph-id", "id": pid},
                locator={"by": "paragraph-id", "id": pid},
                confidence=0.6, source="heuristic:underscore")
    return slots


def extract(path: str | Path) -> dict:
    doc = load_docx(path)
    index = DocIndex(doc)
    paragraphs = [_paragraph_facts(doc, pid, p) for pid, p in index.paragraphs.items()]
    return {
        "schema": FACTS_SCHEMA,
        "template": {
            "path": str(path),
            "content_sha256": content_fingerprint(path),
            "page": _page_info(doc),
        },
        "paragraphs": paragraphs,
        "tables": _table_facts(index),
        "slots": _detect_slots(index),
        "figures": {"count": image_count(doc)},
        "stats": {
            "paragraphs": len(paragraphs),
            "tables": len({t["id"] for t in _table_facts(index)}),
            "characters": sum(len(p["text"]) for p in paragraphs),
        },
    }
