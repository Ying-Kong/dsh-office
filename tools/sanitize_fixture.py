#!/usr/bin/env python3
"""Anonymise a fixture template in place.

Design notes
------------
* Replacements are supplied on the command line (or via ``--map``), never
  hard-coded here: the repository must not embed the very names it redacts.
* Only exact tokens that live inside a *single* ``w:t`` are replaced. Rewriting
  whole paragraphs would collapse run formatting and can change line breaking,
  which would make the fixture stop representing the real form.
* A token split across runs is **reported** (exit 1) instead of silently
  skipped: that means the fixture changed shape and the sanitizer needs review.

Usage::

    python tools/sanitize_fixture.py fixtures/templates/report.docx \
        --replace "Real Name=Redacted Name" \
        --replace "<real-id>=2023000000001" --dry-run

    # or, for a whole form at once:
    python tools/sanitize_fixture.py fixtures/templates/report.docx \
        --map private/redactions.json          # never committed

``--map`` expects a JSON object: ``{"original": "replacement", ...}``.
"""

from __future__ import annotations

import argparse
import json
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def parse_replacements(pairs: list[str] | None, map_path: Path | None) -> dict[str, str]:
    replacements: dict[str, str] = {}
    if map_path:
        replacements.update(json.loads(map_path.read_text(encoding="utf-8")))
    for pair in pairs or []:
        if "=" not in pair:
            raise SystemExit(f"--replace expects old=new, got {pair!r}")
        old, new = pair.split("=", 1)
        replacements[old] = new
    if not replacements:
        raise SystemExit("nothing to do: pass --replace old=new (repeatable) or --map file.json")
    for old, new in replacements.items():
        if not old:
            raise SystemExit("empty search token in replacements")
        if old == new:
            raise SystemExit(f"replacement for {old!r} is a no-op")
    return replacements


def _all_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def sanitize(docx_path: Path, replacements: dict[str, str], dry_run: bool = False) -> dict:
    doc = Document(str(docx_path))
    counts: Counter[str] = Counter()

    for paragraph in _all_paragraphs(doc):
        for node in paragraph._p.iter(qn("w:t")):
            if not node.text:
                continue
            text = node.text
            for needle, pseudonym in replacements.items():
                if needle in text:
                    counts[needle] += text.count(needle)
                    text = text.replace(needle, pseudonym)
            node.text = text

    leftover: Counter[str] = Counter()
    for paragraph in _all_paragraphs(doc):
        for needle in replacements:
            leftover[needle] += paragraph.text.count(needle)

    if not dry_run and counts:
        doc.save(str(docx_path))

    return {"replaced": dict(counts), "leftover": {k: v for k, v in leftover.items() if v}, "dry_run": dry_run}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("document", type=Path)
    parser.add_argument("--replace", action="append", metavar="OLD=NEW")
    parser.add_argument("--map", type=Path, dest="map_path")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--check", action="store_true", help="Report leftovers only; never write.")
    args = parser.parse_args()

    if not zipfile.is_zipfile(args.document):
        raise SystemExit(f"not a .docx package: {args.document}")

    replacements = parse_replacements(args.replace, args.map_path)
    result = sanitize(args.document, replacements, dry_run=args.dry_run or args.check)

    print(f"replaced: {result['replaced'] or '{}'}")
    if result["leftover"]:
        print(f"LEFT OVER (split across runs, needs review): {result['leftover']}")
        raise SystemExit(1)
    if result["dry_run"]:
        print("dry run: file not written")


if __name__ == "__main__":
    main()
